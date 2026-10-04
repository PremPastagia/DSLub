from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re
doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.7)
section.left_margin=section.right_margin=Inches(.8)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.line_spacing=1.12
for n in ['Title','Heading 1','Heading 2']:
 doc.styles[n].font.name='Calibri';doc.styles[n].font.color.rgb=RGBColor(0,0,0)
doc.styles['Title'].font.size=Pt(24)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(18)
footer=section.footer.paragraphs[0];footer.alignment=2
footer.add_run('DS Lub  |  Company evidence checklist  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
fenced=False
for line in Path('docs/Company-Data-Checklist.md').read_text(encoding='utf-8').splitlines():
 if line.startswith('```'):fenced=not fenced;continue
 if not line.strip():continue
 if line.startswith('# '):doc.add_paragraph(line[2:],style='Title');continue
 if line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 1');continue
 line=re.sub(r'\*\*([^*]+)\*\*',r'\1',line).replace('`','')
 if line.startswith('- [ ] '):
  p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.first_line_indent=Inches(-.2)
  r=p.add_run('☐ ');r.font.name='Segoe UI Symbol';p.add_run(line[6:]);continue
 p=doc.add_paragraph(line)
 if fenced:
  for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)
doc.save('docs/DS-Lub-Company-Data-Checklist.docx')
print('Company checklist DOCX created')
